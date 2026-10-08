from __future__ import annotations

import httpx
from openpyxl import load_workbook
from pptx import Presentation

from conftest import make_settings, public_resolver
from nsmpa.fetch import HardenedFetcher
from nsmpa.packet import build_packet
from nsmpa.research import research_all, start_research
from nsmpa.seeds import run_seeds
import test_pipeline as tp
import test_practice as pr
import test_seeds as ts


async def populate(tmp_path, db):
    s = make_settings(tmp_path, research_crime_article_sample=5)
    tp.seed_newsrooms(tmp_path, db)
    await tp.run_research(db, s, tp.FakeSearch(tp.SEARCH_ROUTES))
    p = tmp_path / "s.csv"
    p.write_text("name,url\nCampus Daily,https://daily.example/\n", encoding="utf-8")
    from nsmpa.research import import_entities_csv
    import_entities_csv(db, p, "student_media", "t")
    rid, _ = start_research(db, s, cohort="student_media", limit=None, run_id=None, max_searches=None, command="t")
    f = HardenedFetcher(s, transport=httpx.MockTransport(pr.handler), resolver=public_resolver)
    await research_all(db, s, rid, cohort="student_media", quiet=True, provider=tp.FakeSearch({}), fetcher=f)
    for kind in ("precedents", "experts"):
        await run_seeds(db, s, kind, quiet=True, provider=tp.FakeSearch(ts.ROUTES), fetcher=ts.fetcher(s))
    return s


async def test_packet_workbook_and_deck(tmp_path, db):
    s = await populate(tmp_path, db)
    res = build_packet(db, s, s.output_dir)
    wb = load_workbook(res["workbook"])
    for name in ("Read Me", "Summary", "Closest to My Case", "De-index vs Unpublish", "Named Precedents", "Professional Guidance",
                 "Expert Voices", "Documented Practice", "Opposing Evidence", "Entities", "All Evidence", "Sources", "Search Ledger"):
        assert name in wb.sheetnames, name
    opp = wb["Opposing Evidence"]
    assert opp.max_row > 3  # opposing evidence is present, not filtered out
    links = [c.hyperlink.target for row in wb["All Evidence"].iter_rows() for c in row if c.hyperlink]
    assert links and all(u.startswith("https://") for u in links)
    summary = {r[0].value: r for r in wb["Summary"].iter_rows(min_row=4)}
    assert summary["Student media"][1].value.startswith("PRELIMINARY")
    assert summary["Student media"][5].value == "withheld"  # no percentages without validation
    deck = Presentation(res["presentation"])
    assert res["slides"] == len(deck.slides) >= 8
    alltext = " ".join(sh.text_frame.text for sl in deck.slides for sh in sl.shapes if sh.has_text_frame)
    assert "UNVERIFIED" in alltext and "historical record" in alltext
    assert any(sh.has_chart for sl in deck.slides for sh in sl.shapes)



def test_validate_summary_drops_uncited_and_invented_numbers():
    from nsmpa.packet import validate_summary
    payload = "Evidence:\n[E5] (A, x, adverse) We do not remove stories.\n[E9] (B, y, supportive) Editors may remove names after 2 years."
    kept, dropped = validate_summary([
        {"text": "One newsroom refuses removal.", "evidence_ids": [5]},
        {"text": "Another may remove names after 2 years.", "evidence_ids": [9]},
        {"text": "Most newsrooms (73%) agree.", "evidence_ids": [5]},      # invented number
        {"text": "Unsupported claim.", "evidence_ids": []},                 # no citation
        {"text": "Cites a missing item.", "evidence_ids": [404]},          # unknown id
    ], {5, 9}, payload)
    assert [k["evidence_ids"] for k in kept] == [[5], [9]] and dropped == 3


async def test_packet_with_ai_summaries(tmp_path, db):
    import json as _json
    from types import SimpleNamespace
    s = await populate(tmp_path, db)

    class FakeSummaryAI:
        def __init__(self):
            self.models = SimpleNamespace(generate_content=self.gen)

        def gen(self, *, model, contents, config):
            from google.genai import types
            import re as _re
            ids = [int(x) for x in _re.findall(r"\[E(\d+)\]", contents)][:2]
            body = {"sentences": [{"text": "The cited evidence describes the position.", "evidence_ids": ids},
                                  {"text": "Ninety-nine percent agree (99%).", "evidence_ids": ids}]}
            return SimpleNamespace(text=_json.dumps(body), prompt_feedback=None,
                                   candidates=[SimpleNamespace(finish_reason=types.FinishReason.STOP)],
                                   usage_metadata=SimpleNamespace(prompt_token_count=1, candidates_token_count=1))

    res = build_packet(db, s, s.output_dir, ai_summaries_on=True, ai_client=FakeSummaryAI())
    wb = load_workbook(res["workbook"])
    assert "AI Summaries" in wb.sheetnames
    sentences = [c.value for c in wb["AI Summaries"]["B"]][3:]
    assert sentences and all("99%" not in (x or "") for x in sentences)
    deck = Presentation(res["presentation"])
    assert any("AI-DRAFTED SUMMARY" in sh.text_frame.text for sl in deck.slides for sh in sl.shapes if sh.has_text_frame)
