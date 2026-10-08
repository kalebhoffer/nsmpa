"""Wayback comparison, AI second opinion and accuracy audit (no network, no credits, no API calls)."""
from __future__ import annotations

import csv
import json
from types import SimpleNamespace

import httpx
import test_pipeline as tp
import test_practice as pr
from conftest import make_settings, public_resolver
from openpyxl import load_workbook

from nsmpa.ai_review import run_ai_review, verify_quote
from nsmpa.audit import cohen_kappa, create_sample, export_csv, import_csv, report, wilson
from nsmpa.fetch import HardenedFetcher
from nsmpa.packet import build_packet
from nsmpa.research import import_entities_csv, research_all, start_research
from nsmpa.utils import normalize_for_hash
from nsmpa.wayback import compare_versions, person_names

OLD = (b"<html><head><title>Student John Smithers arrested for assault</title></head><body><main>"
       b"<p>Police arrested John Smithers, a junior, on Tuesday near campus. Smithers was charged with assault.</p>"
       b"</main></body></html>")
NEW_NOINDEX = (b"<html><head><title>Student arrested for assault</title><meta name='robots' content='noindex'></head><body><main>"
               b"<p>Editor's note: charges were dismissed. This story has been updated to remove the student's name.</p>"
               b"<p>Police arrested a junior on Tuesday near campus. The student was charged with assault.</p></main></body></html>")


# ============================================================================ Wayback

def test_compare_versions_detects_changes_without_storing_names():
    obs = compare_versions(OLD, NEW_NOINDEX, url="https://daily.example/a", current_status=200, threshold=0.85)
    types = {o["type"] for o in obs}
    assert {"names_removed", "noindex_added", "title_changed"} <= types
    blob = json.dumps(obs)
    assert "Smithers" not in blob and "John" not in blob  # privacy: names never stored
    nr = next(o for o in obs if o["type"] == "names_removed")
    assert nr["count"] >= 1 and nr["update_note_present"] is True
    assert compare_versions(OLD, None, url="u", current_status=410, threshold=0.85) == [{"type": "unpublished", "current_status": 410}]
    assert compare_versions(OLD, OLD, url="u", current_status=200, threshold=0.85) == []


def test_person_names_skips_institutions():
    names = person_names("John Smithers met Police Chief Brown at Washington State University and Mary Ann O'Neil.")
    assert "John Smithers" in names and not any("University" in n for n in names)


CRIME_URL = "https://daily.example/2014/03/student-arrested-for-assault-near-campus/"


def archive_handler(req: httpx.Request) -> httpx.Response:
    if req.url.host == "web.archive.org":
        if req.url.path == "/robots.txt":
            return httpx.Response(404)
        if req.url.path.startswith("/cdx/"):
            if "student-arrested" in str(req.url):
                body = [["timestamp", "statuscode", "digest"], ["20140305000000", "200", "A"], ["20200101000000", "200", "B"]]
            else:
                body = [["timestamp", "statuscode", "digest"]]
            return httpx.Response(200, headers={"content-type": "application/json"}, content=json.dumps(body).encode())
        if req.url.path.startswith("/web/20140305000000id_/"):
            return httpx.Response(200, headers={"content-type": "text/html"}, content=OLD)
        return httpx.Response(404)
    if str(req.url) == CRIME_URL:
        return httpx.Response(200, headers={"content-type": "text/html"}, content=NEW_NOINDEX)
    return pr.handler(req)


async def test_wayback_in_research_records_archive_practice(tmp_path, db):
    p = tmp_path / "e.csv"
    p.write_text("name,url\nCampus Daily,https://daily.example/\n", encoding="utf-8")
    import_entities_csv(db, p, "student_media", "t")
    s = make_settings(tmp_path, research_crime_article_sample=5)
    rid, _ = start_research(db, s, cohort="student_media", limit=None, run_id=None, max_searches=None, command="t")
    f = HardenedFetcher(s, transport=httpx.MockTransport(archive_handler), resolver=public_resolver)
    await research_all(db, s, rid, cohort="student_media", quiet=True, provider=tp.FakeSearch({}), fetcher=f)
    chk = db.execute("SELECT * FROM wayback_checks WHERE url=?", (CRIME_URL,)).fetchone()
    assert chk["status"] == "changed" and chk["earliest_ts"] == "20140305000000"
    ev = db.execute("SELECT * FROM evidence_items WHERE statement_type='wayback_names_removed'").fetchone()
    assert ev and "Smithers" not in ev["excerpt"] and "Names withheld" in ev["excerpt"]
    assert db.scalar("SELECT COUNT(*) FROM evidence_items WHERE statement_type='wayback_noindex_added'") == 1
    assert db.scalar("SELECT COUNT(*) FROM evidence_items WHERE excerpt LIKE '%Smithers%' OR context LIKE '%Smithers%'") == 0
    st = db.execute("SELECT action_positions_json FROM entity_stances WHERE run_id=?", (rid,)).fetchone()
    assert json.loads(st["action_positions_json"])["anonymize"]["technical"] == "archive_change_observed"


# ============================================================================ AI second opinion

def canned_findings(page: str) -> list[dict]:
    findings = []
    if "We do not remove stories" in page:
        findings = [
            dict(quote="We do not remove stories from our archive.", kind="policy", action="unpublish",
                 position="rejected", direction="adverse", conditions="", speaker="", speaker_role=""),
            dict(quote="We happily delete any story on request.", kind="policy", action="unpublish",  # fabricated
                 position="permitted", direction="supportive", conditions="", speaker="", speaker_role=""),
        ]
    if "Relief Daily removed the name" in page:
        findings = [dict(quote="The Relief Daily removed the name of a former student from a 2015 arrest story after the charges were dismissed",
                         kind="practice", action="anonymize", position="practiced", direction="supportive", conditions="",
                         speaker="Pat Columnist", speaker_role="media columnist"),
                    dict(quote="Another paper in Ohio refused to change anything", kind="opinion", action="none",
                         position="mentioned", direction="neutral", conditions="", speaker="Ohio", speaker_role="")]
    if "served the valley since 1901" in page:
        findings = [dict(quote="The Strict Times has served the valley since 1901", kind="other", action="none",
                         position="mentioned", direction="adverse", conditions="", speaker="", speaker_role="")]
    return findings


class FakeAI:
    """Stands in for google.genai.Client(): client.models.generate_content(model, contents, config)."""
    def __init__(self, finish: str = "STOP", block: str | None = None, fail_codes: list[int] | None = None):
        self.calls = 0
        self.finish, self.block, self.fail_codes = finish, block, list(fail_codes or [])
        self.models = SimpleNamespace(generate_content=self.generate_content)

    def generate_content(self, *, model, contents, config):
        self.calls += 1
        if self.fail_codes:
            err = RuntimeError("transient")
            err.code = self.fail_codes.pop(0)
            raise err
        assert config.response_mime_type == "application/json" and config.response_json_schema["type"] == "object"
        assert "my case" not in config.system_instruction.lower()  # neutral prompt: requester's goal never sent
        findings = canned_findings(contents)
        from google.genai import types
        return SimpleNamespace(
            text=json.dumps({"relevant": bool(findings), "summary": "s", "findings": findings}),
            prompt_feedback=SimpleNamespace(block_reason=self.block) if self.block else None,
            candidates=[SimpleNamespace(finish_reason=getattr(types.FinishReason, self.finish))],
            usage_metadata=SimpleNamespace(prompt_token_count=100, candidates_token_count=50))


class FakeClaude:
    """Stands in for anthropic.Anthropic() when ai_provider='anthropic'."""
    def __init__(self):
        self.calls = 0
        self.beta = SimpleNamespace(messages=SimpleNamespace(create=self.create))
        self.messages = SimpleNamespace(create=self.create)

    def create(self, **kw):
        self.calls += 1
        findings = canned_findings(kw["messages"][0]["content"])
        return SimpleNamespace(stop_reason="end_turn", model=kw["model"], usage=SimpleNamespace(input_tokens=100, output_tokens=50),
                               content=[SimpleNamespace(type="text", text=json.dumps({"relevant": bool(findings), "summary": "s",
                                                                                     "findings": findings}))])


async def test_ai_review_verifies_quotes_compares_and_caches(tmp_path, db):
    s = make_settings(tmp_path)
    tp.seed_newsrooms(tmp_path, db)
    rid, _ = await tp.run_research(db, s, tp.FakeSearch(tp.SEARCH_ROUTES))
    ai = FakeAI()
    res = await run_ai_review(db, s, run_id=rid, quiet=True, client=ai)
    assert res["status"] == "completed" and res["live_calls"] == ai.calls > 0
    rows = {r["quote"]: r for r in db.execute("SELECT * FROM ai_findings")}
    assert rows["We do not remove stories from our archive."]["quote_verified"] == 1
    assert rows["We do not remove stories from our archive."]["agreement"] == "agree"
    assert rows["We happily delete any story on request."]["quote_verified"] == 0  # fabricated quote rejected
    assert rows["We happily delete any story on request."]["agreement"] == "unverified_quote"
    assert db.scalar("SELECT COUNT(*) FROM review_queue WHERE reasons_json LIKE '%ai_rule%'") >= 1
    # Stances are untouched by AI output.
    assert tp.stances(db, rid)["Strict Times"] == "UPDATE_ONLY"
    v = db.execute("SELECT * FROM voices WHERE person_name='Pat Columnist'").fetchone()
    assert v and v["attribution_method"].startswith("ai:") and v["role"] == "media columnist" and v["direction"] == "supportive"
    assert db.scalar("SELECT COUNT(*) FROM voices WHERE person_name='Ohio'") == 0  # single-word "speaker" rejected
    ai2 = FakeAI()
    res2 = await run_ai_review(db, s, run_id=rid, quiet=True, client=ai2)
    assert ai2.calls == 0 and res2["cached"] == res["live_calls"]


async def test_ai_review_call_cap(tmp_path, db):
    s = make_settings(tmp_path)
    tp.seed_newsrooms(tmp_path, db)
    rid, _ = await tp.run_research(db, s, tp.FakeSearch(tp.SEARCH_ROUTES))
    ai = FakeAI()
    res = await run_ai_review(db, s, run_id=rid, quiet=True, client=ai, max_calls=1)
    assert ai.calls <= 1 and res["status"] == "budget_exhausted" and res["remaining"] > 0


async def test_gemini_blocks_and_truncation_are_failures_not_findings(tmp_path, db):
    s = make_settings(tmp_path)
    tp.seed_newsrooms(tmp_path, db)
    rid, _ = await tp.run_research(db, s, tp.FakeSearch(tp.SEARCH_ROUTES))
    for fake in (FakeAI(finish="SAFETY"), FakeAI(finish="MAX_TOKENS"), FakeAI(block="PROHIBITED_CONTENT")):
        db.execute("DELETE FROM ai_cache")
        res = await run_ai_review(db, s, run_id=rid, quiet=True, client=fake)
        assert res["reviewed"] == 0 and res["failed"] > 0
    assert db.scalar("SELECT COUNT(*) FROM ai_findings") == 0


async def test_gemini_retries_transient_errors(tmp_path, db):
    s = make_settings(tmp_path, ai_max_retries=2)
    tp.seed_newsrooms(tmp_path, db)
    rid, _ = await tp.run_research(db, s, tp.FakeSearch(tp.SEARCH_ROUTES))
    import nsmpa.ai_review as air
    air.time.sleep = lambda *_: None  # no real backoff in tests
    fake = FakeAI(fail_codes=[503, 429])
    res = await run_ai_review(db, s, run_id=rid, quiet=True, client=fake, limit=1)
    assert res["reviewed"] == 1 and fake.calls == 3


async def test_anthropic_provider_still_supported(tmp_path, db):
    s = make_settings(tmp_path, ai_provider="anthropic", ai_model="claude-opus-5-5")
    tp.seed_newsrooms(tmp_path, db)
    rid, _ = await tp.run_research(db, s, tp.FakeSearch(tp.SEARCH_ROUTES))
    res = await run_ai_review(db, s, run_id=rid, quiet=True, client=FakeClaude())
    assert res["status"] == "completed" and res["provider"] == "anthropic" and res["verified"] >= 1


def test_missing_gemini_key_is_a_clear_error(tmp_path, monkeypatch):
    from nsmpa.ai_review import AIUnavailable, make_client
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    monkeypatch.delenv("GOOGLE_API_KEY", raising=False)
    try:
        make_client(make_settings(tmp_path))
        raise AssertionError("expected AIUnavailable")
    except AIUnavailable as exc:
        assert "GEMINI_API_KEY" in str(exc)


def test_verify_quote_is_strict_but_whitespace_tolerant():
    page = normalize_for_hash("We do  not remove\nstories from our archive.")
    assert verify_quote("We do not remove stories from our archive", page)
    assert not verify_quote("We do not delete stories", page)
    assert not verify_quote("We do", page)  # too short to be meaningful


# ============================================================================ audit

def test_wilson_and_kappa():
    p, lo, hi = wilson(45, 50)
    assert p == 0.9 and lo < 0.9 < hi and lo > 0.75
    assert wilson(0, 0) is None
    assert cohen_kappa([("a", "a"), ("b", "b"), ("a", "b"), ("b", "b")]) == 0.5
    assert cohen_kappa([("a", "a")] * 5) == 1.0


async def test_audit_sample_label_report(tmp_path, db):
    s = make_settings(tmp_path)
    tp.seed_newsrooms(tmp_path, db)
    await tp.run_research(db, s, tp.FakeSearch(tp.SEARCH_ROUTES))
    smp = create_sample(db, 10, seed=7)
    assert smp["sample_size"] > 0 and set(smp["strata"]) <= {"supportive", "adverse", "neutral"}
    assert create_sample(db, 10, seed=7)["strata"] == smp["strata"]  # reproducible
    out = tmp_path / "audit.csv"
    export_csv(db, smp["audit_id"], out)
    rows = list(csv.DictReader(open(out, encoding="utf-8")))
    assert "machine_direction" not in rows[0]  # blind by default
    machine = {r["audit_item_id"]: r for r in db.execute(
        "SELECT ai.id AS audit_item_id, e.direction FROM audit_items ai JOIN evidence_items e ON e.id=ai.evidence_id WHERE ai.audit_id=?",
        (smp["audit_id"],))}
    for labeler, flip in (("alice", False), ("bob", True)):
        p = tmp_path / f"{labeler}.csv"
        with open(p, "w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
            w.writeheader()
            for i, r in enumerate(rows):
                d = machine[int(r["audit_item_id"])]["direction"]
                if flip and i == 0:
                    d = "neutral" if d != "neutral" else "adverse"
                r = dict(r, **{"relevant (y/n)": "y", "direction (supportive/adverse/neutral)": d})
                w.writerow(r)
        import_csv(db, p, labeler)
    rep = report(db, smp["audit_id"])
    assert rep["direction_accuracy"][0] == 1.0  # alice agreed with every machine label
    assert rep["inter_rater"]["items"] == len(rows) and rep["inter_rater"]["cohen_kappa_direction"] < 1.0
    res = build_packet(db, s, s.output_dir)
    wb = load_workbook(res["workbook"])
    assert {"Archive Changes", "AI Second Opinion", "Accuracy"} <= set(wb.sheetnames)
    assert "Direction accuracy" in [c.value for c in wb["Accuracy"]["A"]]


# ============================================================================ v0.6 stage 1: heartbeats, notify

async def test_research_writes_heartbeat_with_entity_steps(tmp_path, db):
    s = make_settings(tmp_path)
    tp.seed_newsrooms(tmp_path, db)
    rid, _ = await tp.run_research(db, s, tp.FakeSearch(tp.SEARCH_ROUTES))
    hb = db.execute("SELECT * FROM run_heartbeats WHERE run_id=?", (rid,)).fetchone()
    assert hb and hb["finished"] == 1
    state = json.loads(hb["state_json"])
    assert state["completed"] == 4 and state["step_total"] == 8 and state["pages_fetched"] > 0
    from rich.console import Console

    from nsmpa.progress import HeartbeatView
    c = Console(record=True, width=120)
    c.print(HeartbeatView(state, hb["title"]).render())
    out = c.export_text()
    assert "4 / 4" in out and "Entity steps" in out


def test_notify_escapes_and_is_platform_safe(monkeypatch):
    import nsmpa.notify as n
    calls = []
    monkeypatch.setattr(n.subprocess, "run", lambda args, **kw: calls.append(args))
    monkeypatch.setattr(n.sys, "platform", "darwin")
    assert n.notify('Done "now"', "x", enabled=True) and '\\"now\\"' in calls[0][2]
    assert n.notify("t", "m", enabled=False) is False
    monkeypatch.setattr(n.sys, "platform", "linux")
    assert n.notify("t", "m") is False



# ============================================================================ v0.6 stage 5: AI discovery assist

class FakeDiscoveryAI:
    def __init__(self, answer):
        self.answer, self.calls = answer, 0
        self.models = SimpleNamespace(generate_content=self.generate_content)

    def generate_content(self, *, model, contents, config):
        self.calls += 1
        from google.genai import types
        return SimpleNamespace(text=json.dumps(self.answer), prompt_feedback=None,
                               candidates=[SimpleNamespace(finish_reason=types.FinishReason.STOP)],
                               usage_metadata=SimpleNamespace(prompt_token_count=10, candidates_token_count=5))


def _ambiguous_institution(db):
    db.execute("INSERT INTO institutions(unitid,name,website,control,level,included) VALUES('900009','Ambig U','https://ambig.edu/',1,1,1)")
    for url, dom, title in (("https://paper-a.example/", "paper-a.example", "The Ambig Times | student newspaper"),
                            ("https://paper-b.example/", "paper-b.example", "Ambig Review | student news")):
        db.execute("INSERT INTO publication_candidates(unitid,url,domain,title,source,score,verified_score) VALUES(?,?,?,?,?,?,?)",
                   ("900009", url, dom, title, "serper", 0.70, 0.70))
    db.conn.commit()


def test_ai_discovery_pick_breaks_tie_and_is_reviewed(tmp_path, db):
    from nsmpa.ai_discovery import run_ai_discovery
    from nsmpa.discovery import promote_candidates
    s = make_settings(tmp_path)
    _ambiguous_institution(db)
    res = run_ai_discovery(db, s, client=FakeDiscoveryAI({"choice": 2, "confidence": 0.9, "publication_name": "Ambig Review", "reason": "r"}))
    assert res["picked"] == 1
    out = promote_candidates(db, 0.6, use_ai=True)
    pub = db.execute("SELECT * FROM publications WHERE unitid='900009'").fetchone()
    assert out["ai_assisted"] == 1 and pub["domain"] == "paper-b.example" and pub["ambiguous"] == 0
    assert "ai_assist" in pub["discovery_method"]
    assert "ai_assisted_identification" in db.scalar("SELECT reasons_json FROM review_queue WHERE item_type='publication'")


def test_ai_discovery_rejects_out_of_range_choice_and_respects_none(tmp_path, db):
    from nsmpa.ai_discovery import run_ai_discovery
    from nsmpa.discovery import promote_candidates
    s = make_settings(tmp_path)
    _ambiguous_institution(db)
    assert run_ai_discovery(db, s, client=FakeDiscoveryAI({"choice": 7, "confidence": 0.99, "publication_name": "x", "reason": "r"}))["invalid"] == 1
    assert db.scalar("SELECT COUNT(*) FROM publication_candidates WHERE ai_pick_json IS NOT NULL") == 0
    db.execute("DELETE FROM ai_cache")
    run_ai_discovery(db, s, client=FakeDiscoveryAI({"choice": 0, "confidence": 0.95, "publication_name": "", "reason": "press pages"}))
    assert promote_candidates(db, 0.6, use_ai=True)["ai_blocked"] == 1
    assert db.scalar("SELECT COUNT(*) FROM publications WHERE unitid='900009'") == 0
