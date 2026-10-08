from __future__ import annotations

from openpyxl import load_workbook

from nsmpa.packet import build_packet
from nsmpa.redact import PLACEHOLDER, URL_SAFE_PLACEHOLDER, Redactor, redact_rows
import test_packet


def test_redacts_private_names_keeps_public_ones():
    r = Redactor({"The Boston Globe", "Brian McGrory", "State University", "Campus Daily"})
    out = r.text("Police arrested John Smithers near State University. Smithers was charged, The Boston Globe reported; "
                 "Brian McGrory said the Campus Daily would update it.")
    assert "Smithers" not in out and "John" not in out and out.count(PLACEHOLDER) == 2
    assert "Boston Globe" in out and "Brian McGrory" in out and "State University" in out and "Campus Daily" in out


def test_speaker_kept_in_own_quote_and_third_parties_redacted():
    r = Redactor()
    row = {"person_name": "Pat Columnist", "quote": "Pat Columnist said the paper removed Jane Doe's name after the charges were dropped."}
    out = redact_rows([row], r)[0]["quote"]
    assert "Pat Columnist" in out and "Jane Doe" not in out


def test_url_with_redacted_name_is_withheld_per_row():
    r = Redactor()
    rows = redact_rows([
        {"excerpt": "John Smithers was arrested.", "source_url": "https://paper.example/2014/03/john-smithers-arrested/"},
        {"excerpt": "Editors may remove names.", "source_url": "https://paper.example/2014/03/john-hall-renovation/"},
        {"excerpt": "x", "archive_url": "https://web.archive.org/web/2014/x"},
    ], r)
    assert rows[0]["source_url"] == URL_SAFE_PLACEHOLDER
    assert rows[1]["source_url"].startswith("https://")      # names from other rows don't leak into URL decisions
    assert rows[2]["archive_url"] == "[withheld in shared copy]"


def test_institutional_phrases_not_redacted():
    r = Redactor()
    text = "The Supreme Court and Washington State University Police Department reviewed the Code of Ethics on Monday."
    assert r.text(text) == text


async def test_redacted_packet_has_no_private_names(tmp_path, db):
    s = await test_packet.populate(tmp_path, db)
    tp_name = "Smithers"
    db.execute("UPDATE evidence_items SET excerpt=excerpt || ' Police said John Smithers was involved.' WHERE id=(SELECT MIN(id) FROM evidence_items)")
    db.conn.commit()
    plain = build_packet(db, s, s.output_dir / "plain")
    red = build_packet(db, s, s.output_dir / "red", redact=True)
    assert red["redacted"] and red["redactions"] >= 1
    def all_text(path):
        wb = load_workbook(path)
        return " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)
    assert tp_name in all_text(plain["workbook"])
    red_text = all_text(red["workbook"])
    assert tp_name not in red_text and "automated redaction" in red_text
    assert "Brian McGrory" in red_text or "McGrory" in red_text  # seeded expert stays named


async def test_redacted_dashboard(tmp_path, db):
    from nsmpa.gui import write_dashboard
    s = await test_packet.populate(tmp_path, db)
    db.execute("UPDATE evidence_items SET excerpt=excerpt || ' Police said John Smithers was involved.' WHERE id=(SELECT MIN(id) FROM evidence_items)")
    db.conn.commit()
    html = write_dashboard(db, s, tmp_path / "d.html", redact=True).read_text()
    assert "Smithers" not in html and '"redacted": true' in html
